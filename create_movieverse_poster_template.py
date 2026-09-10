from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt


def set_text_style(paragraph, *, size=14, bold=False, color=RGBColor(15, 23, 42), align=PP_ALIGN.LEFT):
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    paragraph.alignment = align


def add_box(slide, x, y, w, h, title, body_lines=None, title_fill=RGBColor(37, 99, 235), border=RGBColor(48, 66, 103)):
    outer = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    outer.fill.solid()
    outer.fill.fore_color.rgb = RGBColor(248, 250, 252)
    outer.line.color.rgb = border
    outer.line.width = Pt(1.2)

    header_h = 1.4
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x + 0.15), Cm(y + 0.15), Cm(w - 0.3), Cm(header_h))
    header.fill.solid()
    header.fill.fore_color.rgb = title_fill
    header.line.fill.background()

    tf = header.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    set_text_style(p, size=14, bold=True, color=RGBColor(255, 255, 255), align=PP_ALIGN.CENTER)

    if body_lines:
        body = slide.shapes.add_textbox(Cm(x + 0.5), Cm(y + 1.85), Cm(w - 1.0), Cm(h - 2.2))
        btf = body.text_frame
        btf.word_wrap = True
        btf.clear()
        for idx, line in enumerate(body_lines):
            para = btf.paragraphs[0] if idx == 0 else btf.add_paragraph()
            para.text = line
            set_text_style(para, size=11, color=RGBColor(30, 41, 59))
            para.space_after = Pt(3)


def add_placeholder_image(slide, x, y, w, h, caption, fig_no):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(226, 232, 240)
    box.line.color.rgb = RGBColor(71, 85, 105)
    box.line.width = Pt(1.0)

    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "Screenshot Placeholder"
    set_text_style(p, size=12, bold=True, color=RGBColor(51, 65, 85), align=PP_ALIGN.CENTER)

    p2 = tf.add_paragraph()
    p2.text = "Drag your application image here"
    set_text_style(p2, size=10, color=RGBColor(71, 85, 105), align=PP_ALIGN.CENTER)

    cap = slide.shapes.add_textbox(Cm(x), Cm(y + h + 0.25), Cm(w), Cm(1.0))
    ctf = cap.text_frame
    ctf.clear()
    cp = ctf.paragraphs[0]
    cp.text = f"Figure {fig_no}. {caption}"
    set_text_style(cp, size=10, color=RGBColor(30, 41, 59), align=PP_ALIGN.CENTER)


def main():
    prs = Presentation()
    prs.slide_width = Cm(70)
    prs.slide_height = Cm(100)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0), Cm(0), Cm(70), Cm(100))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bg.line.fill.background()

    # Header bar
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0), Cm(0), Cm(70), Cm(14))
    header.fill.solid()
    header.fill.fore_color.rgb = RGBColor(15, 23, 42)
    header.line.fill.background()

    # Logo placeholders
    for x in (1.2, 60.8):
        logo = slide.shapes.add_shape(MSO_SHAPE.OVAL, Cm(x), Cm(1.6), Cm(8), Cm(8))
        logo.fill.solid()
        logo.fill.fore_color.rgb = RGBColor(241, 245, 249)
        logo.line.color.rgb = RGBColor(37, 99, 235)
        logo.line.width = Pt(1.5)
        ltf = logo.text_frame
        ltf.clear()
        lp = ltf.paragraphs[0]
        lp.text = "UNIVERSITY\nLOGO"
        set_text_style(lp, size=11, bold=True, color=RGBColor(30, 41, 59), align=PP_ALIGN.CENTER)

    # Header text
    title = slide.shapes.add_textbox(Cm(10.5), Cm(1.8), Cm(49), Cm(3.2))
    tf = title.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "MOVIEVERSE"
    set_text_style(p, size=34, bold=True, color=RGBColor(255, 255, 255), align=PP_ALIGN.CENTER)

    subtitle = slide.shapes.add_textbox(Cm(10.5), Cm(5.0), Cm(49), Cm(2.2))
    stf = subtitle.text_frame
    stf.clear()
    sp = stf.paragraphs[0]
    sp.text = "Desktop Movie Recommendation System Powered by Machine Learning"
    set_text_style(sp, size=14, color=RGBColor(191, 219, 254), align=PP_ALIGN.CENTER)

    info = slide.shapes.add_textbox(Cm(10.5), Cm(7.2), Cm(49), Cm(4.3))
    itf = info.text_frame
    itf.clear()
    ip1 = itf.paragraphs[0]
    ip1.text = "Student Name Surname"
    set_text_style(ip1, size=16, bold=True, color=RGBColor(255, 255, 255), align=PP_ALIGN.CENTER)
    ip2 = itf.add_paragraph()
    ip2.text = "Supervisor: Dr. ................................"
    set_text_style(ip2, size=12, color=RGBColor(226, 232, 240), align=PP_ALIGN.CENTER)
    ip3 = itf.add_paragraph()
    ip3.text = "Department of Computer Engineering  |  Academic Year: 2025-2026"
    set_text_style(ip3, size=11, color=RGBColor(203, 213, 225), align=PP_ALIGN.CENTER)

    # Left column sections
    add_box(
        slide, 1.5, 15.5, 32.5, 17.0, "INTRODUCTION",
        [
            "MovieVerse is a production-oriented desktop application for personalized movie discovery.",
            "The system analyzes user behavior and profile preferences to recommend relevant films.",
            "It combines GUI engineering, database workflows, and machine learning-ready data processing.",
        ],
    )
    add_box(
        slide, 1.5, 33.8, 32.5, 14.2, "PURPOSE",
        [
            "• Build a practical personalized recommendation engine.",
            "• Improve movie discovery efficiency and user satisfaction.",
            "• Provide top-1 to top-5 ranked recommendations.",
            "• Establish an extendable architecture for future ML upgrades.",
        ],
    )
    add_box(
        slide, 1.5, 49.2, 32.5, 16.8, "MATERIAL & METHODS",
        [
            "• Python desktop GUI application",
            "• PostgreSQL data management",
            "• Recommendation system logic (profile + interactions)",
            "• Feature engineering and category normalization",
            "• ML-ready training views for offline evaluation",
        ],
    )
    add_box(
        slide, 1.5, 67.1, 32.5, 17.4, "TECHNICAL HIGHLIGHTS",
        [
            "• Profile-based watched history tracking",
            "• Interaction-event logging pipeline",
            "• Missing-poster diagnostics and data quality checks",
            "• Recommendation dialog integrated into core UX",
            "• Production-friendly modular architecture",
        ],
        title_fill=RGBColor(79, 70, 229),
    )

    # Right column sections
    add_box(
        slide, 36.0, 15.5, 32.5, 14.0, "RESULTS",
        [
            "• Personalized recommendation workflow implemented end-to-end.",
            "• User-specific top-K recommendation (1-5) is operational.",
            "• Data schema supports scalable model iteration.",
            "• Training and feature views are reusable for experiments.",
        ],
    )

    # Screenshot placeholders
    add_placeholder_image(slide, 36.6, 30.1, 15.0, 10.0, "Main interface of MovieVerse.", 1)
    add_placeholder_image(slide, 52.9, 30.1, 15.0, 10.0, "Recommendation dialog output.", 2)
    add_placeholder_image(slide, 36.6, 41.7, 31.3, 11.2, "Movie detail and user confirmation screen.", 3)

    add_box(
        slide, 36.0, 54.0, 32.5, 13.8, "ARCHITECTURE",
        [
            "Presentation Layer: Desktop GUI and interaction controls",
            "Data Layer: Movies, profiles, watched history, and events",
            "Intelligence Layer: Feature extraction and recommendation ranking",
            "Output Layer: Personalized top-K candidate list",
        ],
        title_fill=RGBColor(79, 70, 229),
    )
    add_box(
        slide, 36.0, 68.6, 32.5, 12.4, "ARGUMENT",
        [
            "MovieVerse demonstrates that a desktop-first engineering product can deliver",
            "measurable personalization using practical ML-driven recommendation workflows.",
            "The architecture is maintainable, modular, and suitable for future ranking models.",
        ],
        title_fill=RGBColor(30, 64, 175),
    )
    add_box(
        slide, 36.0, 81.8, 32.5, 8.2, "FUTURE WORK",
        [
            "• Advanced ranking models (LightGBM / LightFM)",
            "• Online feedback loop and A/B evaluation",
            "• Diversity and novelty-aware recommendation objectives",
        ],
        title_fill=RGBColor(30, 64, 175),
    )

    # Footer
    footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0), Cm(94), Cm(70), Cm(6))
    footer.fill.solid()
    footer.fill.fore_color.rgb = RGBColor(15, 23, 42)
    footer.line.fill.background()

    footer_text = slide.shapes.add_textbox(Cm(1.2), Cm(95.4), Cm(67.6), Cm(3.6))
    ftf = footer_text.text_frame
    ftf.clear()
    fp = ftf.paragraphs[0]
    fp.text = "University Name | Faculty of Engineering | Department of Computer Engineering | Graduation Project Poster"
    set_text_style(fp, size=11, color=RGBColor(226, 232, 240), align=PP_ALIGN.CENTER)
    fp2 = ftf.add_paragraph()
    fp2.text = "Replace placeholders with your personal information and application screenshots."
    set_text_style(fp2, size=10, color=RGBColor(148, 163, 184), align=PP_ALIGN.CENTER)

    output = "MovieVerse_Poster_Template_70x100.pptx"
    prs.save(output)
    print(output)


if __name__ == "__main__":
    main()
